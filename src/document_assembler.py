"""Assemble the final DOCX report and a professional markdown QC report."""

import logging
import os
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from src.metrics import count_improved_sections, count_resolved_issues
from src.utils import OUTPUT_DIR

logger = logging.getLogger(__name__)

FINAL_REPORT_PATH = OUTPUT_DIR / "Final_Report.docx"
QC_REPORT_PATH = OUTPUT_DIR / "QC_Report.md"

SUBTITLE = "Prepared by AI Document Production System"

INLINE_PATTERN = re.compile(r"(\*\*[^*\n]+?\*\*|\*[^*\n]+?\*|(?<!\w)_[^_\n]+?_(?!\w))")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.*)")
BULLET_PATTERN = re.compile(r"^\s*[-*]\s+(.*)")
NUMBERED_PATTERN = re.compile(r"^\s*\d+[.)]\s+(.*)")
TABLE_SEPARATOR_PATTERN = re.compile(r"^:?\s*-{3,}\s*:?$")

CHARS_PER_PAGE = 2500


class DocumentAssembler:
    """Build the final DOCX report and the QC markdown report in output/."""

    def __init__(self) -> None:
        self.output_dir = OUTPUT_DIR
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def assemble_document(self, sections: dict, outline: dict, qc_report: dict) -> str:
        """Render sections into a styled DOCX and save it to output/Final_Report.docx.

        Returns:
            Path to the saved DOCX file.
        """
        doc = Document()
        self._add_title_page(doc, qc_report)

        outline_sections = outline.get("outline", []) if isinstance(outline, dict) else []
        for idx, section in enumerate(outline_sections):
            title = section.get("title", f"Section {idx + 1}")
            text = sections.get(title, "") if isinstance(sections, dict) else ""
            if not text:
                logger.warning("No content found for section: %s", title)
                placeholder = doc.add_paragraph()
                placeholder.add_run(f"[{title} — content pending]").italic = True
                continue
            self._render_markdown(doc, text)

        doc.save(FINAL_REPORT_PATH)

        char_count = sum(len(t) for t in sections.values())
        pages_estimate = max(1, char_count // CHARS_PER_PAGE)
        size_kb = os.path.getsize(FINAL_REPORT_PATH) / 1024
        logger.info(
            "Assembled document: ~%d pages, %.1f KB -> %s",
            pages_estimate,
            size_kb,
            FINAL_REPORT_PATH,
        )
        return str(FINAL_REPORT_PATH)

    def _add_title_page(self, doc: Document, qc_report: dict) -> None:
        title = os.getenv("REPORT_TITLE", "Market Entry Report")
        for _ in range(6):
            doc.add_paragraph()

        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(24)

        subtitle_p = doc.add_paragraph()
        subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_p.add_run(SUBTITLE).font.size = Pt(14)

        date_p = doc.add_paragraph()
        date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        date_p.add_run(date.today().strftime("%B %d, %Y")).font.size = Pt(12)

        if isinstance(qc_report, dict) and qc_report.get("overall_score") is not None:
            qc_p = doc.add_paragraph()
            qc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            qc_p.add_run(f"AI Quality Score: {qc_report.get('overall_score')}/100").font.size = Pt(11)

        doc.add_page_break()

    def _render_markdown(self, doc: Document, text: str) -> None:
        lines = text.splitlines()
        index = 0
        while index < len(lines):
            stripped = lines[index].strip()
            if not stripped or stripped in ("---", "***", "___"):
                index += 1
                continue
            if stripped.startswith("|"):
                block = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    block.append(lines[index].strip())
                    index += 1
                self._render_table_block(doc, block)
                continue
            heading = HEADING_PATTERN.match(stripped)
            if heading:
                level = min(len(heading.group(1)), 4)
                self._add_heading(doc, heading.group(2).strip(), level)
                index += 1
                continue
            bullet = BULLET_PATTERN.match(stripped)
            if bullet:
                self._add_paragraph_with_bold(doc, bullet.group(1).strip(), style="List Bullet")
                index += 1
                continue
            numbered = NUMBERED_PATTERN.match(stripped)
            if numbered:
                self._add_paragraph_with_bold(doc, numbered.group(1).strip(), style="List Number")
                index += 1
                continue
            self._add_paragraph_with_bold(doc, stripped)
            index += 1

    @staticmethod
    def _split_table_row(line: str) -> list[str]:
        return [cell.strip() for cell in line.strip("|").split("|")]

    @staticmethod
    def _is_separator_row(cells: list[str]) -> bool:
        return bool(cells) and all(TABLE_SEPARATOR_PATTERN.match(cell) for cell in cells)

    def _render_table_block(self, doc: Document, block: list[str]) -> None:
        rows = [cells for cells in (self._split_table_row(line) for line in block) if not self._is_separator_row(cells)]
        if len(rows) < 2:
            if rows:
                self._add_paragraph_with_bold(doc, " | ".join(rows[0]))
            return

        column_count = max(len(row) for row in rows)
        table = doc.add_table(rows=len(rows), cols=column_count)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index in range(column_count):
                value = row[column_index] if column_index < len(row) else ""
                paragraph = table.cell(row_index, column_index).paragraphs[0]
                self._add_inline_runs(paragraph, value)

    def _add_heading(self, doc: Document, title: str, level: int) -> None:
        heading = doc.add_heading("", level=level)
        run = heading.add_run(title.replace("**", ""))
        run.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)

    def _add_paragraph_with_bold(self, doc: Document, text: str, style: str | None = None) -> None:
        paragraph = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        self._add_inline_runs(paragraph, text)

    @staticmethod
    def _add_inline_runs(paragraph, text: str) -> None:
        position = 0
        for match in INLINE_PATTERN.finditer(text):
            if match.start() > position:
                paragraph.add_run(text[position : match.start()])
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("*"):
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            position = match.end()
        if position < len(text):
            paragraph.add_run(text[position:])

    @staticmethod
    def _count_by_severity(issues: list, severity: str) -> int:
        return sum(1 for issue in issues if isinstance(issue, dict) and issue.get("severity") == severity)

    @staticmethod
    def _render_issues_markdown(issues: list, show_none: bool) -> str:
        """Render severity-grouped issue lists; append 'None' blocks when asked."""
        content = ""
        for severity in ("HIGH", "MEDIUM", "LOW"):
            severity_issues = [
                issue for issue in issues if isinstance(issue, dict) and issue.get("severity") == severity
            ]
            if not severity_issues:
                if show_none:
                    content += f"### {severity} Severity: None\n\n"
                continue
            content += f"### {severity} Severity ({len(severity_issues)} issues)\n\n"
            for idx, issue in enumerate(severity_issues, start=1):
                content += f"{idx}. **{issue.get('section', 'N/A')}**: {issue.get('description', 'N/A')}\n"
                content += f"   - Location: {issue.get('location', 'N/A')}\n"
                content += f"   - Suggestion: {issue.get('suggestion', 'N/A')}\n\n"
        return content

    @staticmethod
    def _normalize_qc_reports(qc_report_draft: dict, qc_report_final: dict) -> tuple[dict, dict]:
        """Validate report shapes and default missing issue lists to empty."""
        if not isinstance(qc_report_draft, dict):
            logger.warning("Draft QC report is not a dict; using empty draft")
            qc_report_draft = {}
        if not isinstance(qc_report_final, dict):
            logger.warning("Final QC report is not a dict; using empty final")
            qc_report_final = {}
        if not isinstance(qc_report_draft.get("issues"), list):
            qc_report_draft["issues"] = []
        if not isinstance(qc_report_final.get("issues"), list):
            qc_report_final["issues"] = []
        return qc_report_draft, qc_report_final

    @staticmethod
    def _quality_metrics(qc_report_draft: dict, qc_report_final: dict) -> dict:
        """Compute before/after quality metrics for the QC report."""
        draft_score = qc_report_draft.get("overall_score", 0)
        final_score = qc_report_final.get("overall_score", 0)
        draft_issues = len(qc_report_draft["issues"])
        final_issues = len(qc_report_final["issues"])
        resolved = count_resolved_issues(qc_report_draft["issues"], qc_report_final["issues"])
        resolved_sections = count_improved_sections(qc_report_draft["issues"], qc_report_final["issues"])
        score_delta = final_score - draft_score
        score_delta_text = f"+{score_delta}" if score_delta >= 0 else str(score_delta)
        resolution_pct = int(resolved / draft_issues * 100) if draft_issues > 0 else 0
        return {
            "draft_score": draft_score,
            "final_score": final_score,
            "draft_issues": draft_issues,
            "final_issues": final_issues,
            "resolved": resolved,
            "resolved_sections": resolved_sections,
            "score_delta_text": score_delta_text,
            "resolution_pct": resolution_pct,
        }

    @staticmethod
    def _revision_actions(issues: list, resolved: int, draft_issues: int) -> str:
        """Render the revision-actions block of the QC report."""
        remaining_high = sum(1 for i in issues if isinstance(i, dict) and i.get("severity") == "HIGH")
        if remaining_high == 0:
            content = "- All HIGH-severity issues from draft were addressed by AI Revision Loop\n"
        else:
            content = (
                f"- {remaining_high} HIGH-severity issue(s) remain after AI revision — flagged below for human review\n"
            )
        return content + f"- {resolved} out of {draft_issues} issues resolved automatically\n"

    def _build_quality_report_content(self, qc_report_draft: dict, qc_report_final: dict) -> str:
        """Render the QC report markdown from validated draft and final reviews."""
        m = self._quality_metrics(qc_report_draft, qc_report_final)
        content = f"""# Quality Control Report

## Summary

### Draft Quality (before AI revision)
- Score: {m["draft_score"]}/100
- Total Issues: {m["draft_issues"]}
  - HIGH: {self._count_by_severity(qc_report_draft["issues"], "HIGH")}
  - MEDIUM: {self._count_by_severity(qc_report_draft["issues"], "MEDIUM")}
  - LOW: {self._count_by_severity(qc_report_draft["issues"], "LOW")}

### Final Quality (after AI revision)
- Score: {m["final_score"]}/100
- Total Issues: {m["final_issues"]}
  - HIGH: {self._count_by_severity(qc_report_final["issues"], "HIGH")}
  - MEDIUM: {self._count_by_severity(qc_report_final["issues"], "MEDIUM")}
  - LOW: {self._count_by_severity(qc_report_final["issues"], "LOW")}

### Improvement Metrics
- Score Improvement: {m["draft_score"]} → {m["final_score"]} (**{m["score_delta_text"]}**)
- Issues Resolved: {m["resolved"]}/{m["draft_issues"]} (**{m["resolution_pct"]}%**)
- Sections Improved: {m["resolved_sections"]}

---

## Draft Issues (before revision)

"""
        content += self._render_issues_markdown(qc_report_draft["issues"], show_none=False)
        content += "\n---\n\n## Final Issues (remaining after revision)\n\n"
        content += self._render_issues_markdown(qc_report_final["issues"], show_none=True)
        content += "\n---\n\n## Revision Actions Taken\n\n"
        content += self._revision_actions(qc_report_final["issues"], m["resolved"], m["draft_issues"])
        return content

    def generate_quality_report(self, qc_report_draft: dict, qc_report_final: dict) -> str:
        """Write a QC report with before/after metrics to output/QC_Report.md.

        Args:
            qc_report_draft: QC review of the sections before AI revision.
            qc_report_final: QC review of the sections after AI revision.

        Returns:
            Path to the saved QC report file.
        """
        qc_report_draft, qc_report_final = self._normalize_qc_reports(qc_report_draft, qc_report_final)
        content = self._build_quality_report_content(qc_report_draft, qc_report_final)

        QC_REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
        Path(QC_REPORT_PATH).write_text(content, encoding="utf-8")
        logger.info(
            "QC report with before/after metrics saved to %s (%d bytes)",
            QC_REPORT_PATH,
            QC_REPORT_PATH.stat().st_size,
        )
        return str(QC_REPORT_PATH)
