"""Offline tests for markdown rendering in the DOCX document assembler."""

from docx import Document

from src.document_assembler import DocumentAssembler


def _build_doc(text: str, tmp_path):
    """Render markdown into a DOCX in a temp dir and read it back."""
    assembler = DocumentAssembler()
    doc = Document()
    assembler._render_markdown(doc, text)
    path = tmp_path / "rendered.docx"
    doc.save(path)
    return Document(path)


class TestMarkdownTable:
    def test_table_has_header_and_data_rows_without_separator(self, tmp_path):
        markdown = "| Name | Value |\n| --- | --- |\n| Alpha | 1 |\n| Beta | 2 |\n"
        doc = _build_doc(markdown, tmp_path)

        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert len(table.rows) == 3
        assert len(table.columns) == 2
        assert [cell.text for cell in table.rows[0].cells] == ["Name", "Value"]
        assert [cell.text for cell in table.rows[1].cells] == ["Alpha", "1"]
        assert [cell.text for cell in table.rows[2].cells] == ["Beta", "2"]

    def test_table_cells_have_no_edge_pipes_or_extra_spaces(self, tmp_path):
        markdown = "|  Name  |  Value  |\n| :---: | --- |\n|   Alpha   |   1   |\n|  Beta  |  2  |\n"
        doc = _build_doc(markdown, tmp_path)

        table = doc.tables[0]
        for row in table.rows:
            for cell in row.cells:
                assert "|" not in cell.text
                assert cell.text == cell.text.strip()

    def test_ragged_table_pads_missing_cells(self, tmp_path):
        markdown = "| A | B | C |\n| --- | --- | --- |\n| 1 | 2 |\n"
        doc = _build_doc(markdown, tmp_path)

        table = doc.tables[0]
        assert len(table.columns) == 3
        assert [cell.text for cell in table.rows[1].cells] == ["1", "2", ""]


class TestInlineMarkdown:
    def test_asterisk_italic_is_rendered(self, tmp_path):
        doc = _build_doc("This is *italic* text", tmp_path)

        paragraph = doc.paragraphs[0]
        assert "*" not in paragraph.text
        assert "_" not in paragraph.text
        italic_runs = [run for run in paragraph.runs if run.italic]
        assert [run.text for run in italic_runs] == ["italic"]

    def test_underscore_italic_is_rendered(self, tmp_path):
        doc = _build_doc("This is _italic_ text", tmp_path)

        paragraph = doc.paragraphs[0]
        assert "_" not in paragraph.text
        italic_runs = [run for run in paragraph.runs if run.italic]
        assert [run.text for run in italic_runs] == ["italic"]

    def test_bold_still_rendered(self, tmp_path):
        doc = _build_doc("This is **bold** text", tmp_path)

        paragraph = doc.paragraphs[0]
        assert "**" not in paragraph.text
        bold_runs = [run for run in paragraph.runs if run.bold]
        assert [run.text for run in bold_runs] == ["bold"]


class TestHeadingsAndLists:
    def test_headings_and_bullets_render_as_before(self, tmp_path):
        markdown = "# Title\n\n## Subtitle\n\n- bullet item\n"
        doc = _build_doc(markdown, tmp_path)

        styles = {p.style.name: p.text for p in doc.paragraphs if p.text}
        assert styles.get("Heading 1") == "Title"
        assert styles.get("Heading 2") == "Subtitle"
        assert styles.get("List Bullet") == "bullet item"
